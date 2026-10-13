"""Documentos de oficina para APOLO: crear Word / Excel / PowerPoint / PDF desde un JSON y leer su texto.
uso:
  python scripts/doc.py crear <spec.json>        → crea el archivo (ruta en "salida") e imprime un resumen
  python scripts/doc.py leer <archivo> [max]     → texto de .docx .xlsx .pptx .pdf (max caracteres, 60000 por defecto)
Librerías: python-docx, openpyxl, python-pptx, pypdf, reportlab (pip install --user ...). Ver SKILL.md para el formato."""
import json, os, sys

def falta(mod):
    sys.exit(f'Falta la librería {mod}. Instálala con: python -m pip install --user {mod}')

# ---------------- bloques comunes (Word / PDF) ----------------
# {"h1": t} {"h2": t} {"h3": t} {"p": t} {"lista": [..]} {"numerada": [..]} {"tabla": {"cabecera": [..], "filas": [[..]]}}
# {"imagen": ruta, "ancho_cm": 12} {"cita": t} {"salto": true}
# En los textos, **negrita** y *cursiva*.
import re
def trozos(t):
    """'hola **mundo** y *más*' → [(texto, negrita, cursiva)]"""
    out = []
    for m in re.finditer(r'\*\*(.+?)\*\*|\*(.+?)\*|([^*]+|\*)', str(t)):
        if m.group(1): out.append((m.group(1), True, False))
        elif m.group(2): out.append((m.group(2), False, True))
        else: out.append((m.group(3), False, False))
    return out

def crear_docx(S):
    try:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.enum.text import WD_BREAK
    except ImportError: falta('python-docx')
    d = Document()
    st = d.styles['Normal']; st.font.name = S.get('fuente', 'Calibri'); st.font.size = Pt(S.get('tamano', 11))
    if S.get('titulo'): d.add_heading(S['titulo'], 0)
    def parrafo(t, estilo=None):
        p = d.add_paragraph(style=estilo)
        for txt, b, i in trozos(t): r = p.add_run(txt); r.bold = b; r.italic = i
        return p
    n = {'tablas': 0, 'imagenes': 0, 'parrafos': 0}
    for b in S.get('bloques', []):
        if 'h1' in b: d.add_heading(b['h1'], 1)
        elif 'h2' in b: d.add_heading(b['h2'], 2)
        elif 'h3' in b: d.add_heading(b['h3'], 3)
        elif 'p' in b: parrafo(b['p']); n['parrafos'] += 1
        elif 'cita' in b: parrafo(b['cita'], 'Intense Quote')
        elif 'lista' in b:
            for x in b['lista']: parrafo(x, 'List Bullet')
        elif 'numerada' in b:
            for x in b['numerada']: parrafo(x, 'List Number')
        elif 'tabla' in b:
            T = b['tabla']; cab = T.get('cabecera') or []; filas = T.get('filas') or []
            ncol = max([len(cab)] + [len(f) for f in filas] + [1])
            t = d.add_table(rows=0, cols=ncol); t.style = T.get('estilo', 'Light Grid Accent 1')
            if cab:
                c = t.add_row().cells
                for i, v in enumerate(cab): c[i].text = str(v); c[i].paragraphs[0].runs[0].bold = True
            for f in filas:
                c = t.add_row().cells
                for i, v in enumerate(f): c[i].text = '' if v is None else str(v)
            n['tablas'] += 1
        elif 'imagen' in b:
            d.add_picture(b['imagen'], width=Cm(b.get('ancho_cm', 14))); n['imagenes'] += 1
        elif b.get('salto'): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.save(S['salida']); return n

def crear_pdf(S):
    try:
        from reportlab.lib.pagesizes import A4, letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, ListFlowable, ListItem
    except ImportError: falta('reportlab')
    from xml.sax.saxutils import escape
    ss = getSampleStyleSheet()
    def html(t):
        return ''.join(f'<b>{escape(x)}</b>' if b else f'<i>{escape(x)}</i>' if i else escape(x) for x, b, i in trozos(t))
    doc = SimpleDocTemplate(S['salida'], pagesize=letter if S.get('papel') == 'carta' else A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm, title=S.get('titulo', ''))
    h = []
    if S.get('titulo'): h += [Paragraph(escape(S['titulo']), ss['Title']), Spacer(1, 6)]
    for b in S.get('bloques', []):
        if 'h1' in b: h.append(Paragraph(escape(b['h1']), ss['Heading1']))
        elif 'h2' in b: h.append(Paragraph(escape(b['h2']), ss['Heading2']))
        elif 'h3' in b: h.append(Paragraph(escape(b['h3']), ss['Heading3']))
        elif 'p' in b: h.append(Paragraph(html(b['p']), ss['BodyText']))
        elif 'cita' in b: h.append(Paragraph(html(b['cita']), ParagraphStyle('cita', parent=ss['BodyText'], leftIndent=18, textColor=colors.HexColor('#555555'), fontName='Helvetica-Oblique')))
        elif 'lista' in b or 'numerada' in b:
            items = b.get('lista') or b.get('numerada')
            h.append(ListFlowable([ListItem(Paragraph(html(x), ss['BodyText'])) for x in items], bulletType='1' if 'numerada' in b else 'bullet'))
        elif 'tabla' in b:
            T = b['tabla']; datos = ([T['cabecera']] if T.get('cabecera') else []) + (T.get('filas') or [])
            datos = [[Paragraph(html('' if v is None else v), ss['BodyText']) for v in f] for f in datos]
            t = Table(datos, repeatRows=1 if T.get('cabecera') else 0)
            est = [('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#999999')), ('VALIGN', (0, 0), (-1, -1), 'TOP')]
            if T.get('cabecera'): est.append(('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#dde6f3')))
            t.setStyle(TableStyle(est)); h.append(t)
        elif 'imagen' in b:
            from reportlab.lib.utils import ImageReader
            iw, ih = ImageReader(b['imagen']).getSize(); w = b.get('ancho_cm', 14) * cm
            h.append(Image(b['imagen'], width=w, height=w * ih / iw))
        elif b.get('salto'): h.append(PageBreak())
        h.append(Spacer(1, 6))
    doc.build(h); return {'bloques': len(S.get('bloques', []))}

def crear_xlsx(S):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.chart import BarChart, LineChart, PieChart, Reference
        from openpyxl.utils import get_column_letter
    except ImportError: falta('openpyxl')
    wb = Workbook(); wb.remove(wb.active); res = []
    for H in S.get('hojas', []):
        ws = wb.create_sheet(str(H.get('nombre', 'Hoja'))[:31])
        cab = H.get('cabecera') or []
        if cab:
            ws.append(cab)
            for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='305496')
            ws.freeze_panes = 'A2'
        for f in H.get('filas') or []: ws.append(f)           # textos que empiezan por "=" son fórmulas
        for i, a in enumerate(H.get('anchos') or [], 1): ws.column_dimensions[get_column_letter(i)].width = a
        if not H.get('anchos'):
            for i in range(1, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(i)].width = min(45, max(8, max(len(str(c.value or '')) for c in ws[get_column_letter(i)]) + 2))
        for col, fmt in (H.get('formatos') or {}).items():   # {"C": "#,##0.00 €", "D": "0%"}
            for c in ws[col][1 if cab else 0:]: c.number_format = fmt
        G = H.get('grafico')
        if G and ws.max_row > 1:
            ch = {'lineas': LineChart, 'tarta': PieChart}.get(G.get('tipo'), BarChart)()
            ch.title = G.get('titulo', ''); n = ws.max_row
            for col in G.get('cols_valores', [2]):
                ch.add_data(Reference(ws, min_col=col, min_row=1 if cab else 2, max_row=n), titles_from_data=bool(cab))
            ch.set_categories(Reference(ws, min_col=G.get('col_categorias', 1), min_row=2 if cab else 1, max_row=n))
            ch.width, ch.height = 18, 9
            ws.add_chart(ch, G.get('celda', get_column_letter(ws.max_column + 2) + '2'))
        res.append(f"{ws.title}: {ws.max_row} filas × {ws.max_column} columnas{' + gráfico' if G else ''}")
    wb.save(S['salida']); return {'hojas': res}

def crear_pptx(S):
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError: falta('python-pptx')
    pr = Presentation(); pr.slide_width, pr.slide_height = Inches(13.333), Inches(7.5)   # 16:9
    if S.get('titulo'):
        s = pr.slides.add_slide(pr.slide_layouts[0]); s.shapes.title.text = S['titulo']
        if len(s.placeholders) > 1: s.placeholders[1].text = S.get('subtitulo', '')
    for D in S.get('diapositivas', []):
        s = pr.slides.add_slide(pr.slide_layouts[1 if D.get('puntos') else 5])
        s.shapes.title.text = D.get('titulo', '')
        if D.get('puntos'):
            cuerpo = s.placeholders[1]
            if D.get('imagen'): cuerpo.width = Inches(6.6)
            tf = cuerpo.text_frame; tf.clear()
            for i, x in enumerate(D['puntos']):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                nivel = 1 if str(x).startswith('  ') else 0
                p.text = str(x).strip(); p.level = nivel; p.font.size = Pt(24 if nivel == 0 else 20)
        if D.get('imagen'):
            izq = Inches(7.3) if D.get('puntos') else Inches(1.5)
            s.shapes.add_picture(D['imagen'], izq, Inches(1.7), width=Inches(5.5) if D.get('puntos') else Inches(10.3))
        if D.get('notas'): s.notes_slide.notes_text_frame.text = D['notas']
    pr.save(S['salida']); return {'diapositivas': len(pr.slides)}

# ---------------- leer ----------------
def leer(f, maximo=60000):
    ext = os.path.splitext(f)[1].lower(); out = []
    if ext == '.docx':
        from docx import Document
        d = Document(f)
        for p in d.paragraphs:
            if p.text.strip(): out.append(('#' * int(p.style.name[-1]) + ' ' if p.style.name.startswith('Heading') and p.style.name[-1].isdigit() else '') + p.text)
        for i, t in enumerate(d.tables, 1):
            out.append(f'\n[Tabla {i}]'); out += [' | '.join(c.text for c in r.cells) for r in t.rows]
    elif ext in ('.xlsx', '.xlsm'):
        from openpyxl import load_workbook
        wb = load_workbook(f, data_only=True, read_only=True)
        for ws in wb.worksheets:
            out.append(f'\n[Hoja {ws.title}]')
            for k, r in enumerate(ws.iter_rows(values_only=True)):
                if k >= 500: out.append('… (más filas)'); break
                if any(v is not None for v in r): out.append('\t'.join('' if v is None else str(v) for v in r))
    elif ext == '.pptx':
        from pptx import Presentation
        for i, s in enumerate(Presentation(f).slides, 1):
            out.append(f'\n[Diapositiva {i}]')
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip(): out.append(sh.text_frame.text)
            if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip(): out.append('Notas: ' + s.notes_slide.notes_text_frame.text)
    elif ext == '.pdf':
        from pypdf import PdfReader
        for i, p in enumerate(PdfReader(f).pages, 1):
            out.append(f'\n[Página {i}]\n' + (p.extract_text() or ''))
            if sum(len(x) for x in out) > maximo: break
    else: sys.exit(f'No sé leer {ext} (sí: .docx .xlsx .pptx .pdf)')
    t = '\n'.join(out).strip()
    return t[:maximo] + ('\n… (recortado)' if len(t) > maximo else '')

if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    if len(sys.argv) < 3: sys.exit(__doc__)
    if sys.argv[1] == 'leer':
        print(leer(sys.argv[2], int(sys.argv[3]) if len(sys.argv) > 3 else 60000))
    elif sys.argv[1] == 'crear':
        S = json.load(open(sys.argv[2], encoding='utf-8'))
        S['salida'] = os.path.abspath(S['salida']); os.makedirs(os.path.dirname(S['salida']), exist_ok=True)
        ext = os.path.splitext(S['salida'])[1].lower()
        f = {'.docx': crear_docx, '.pdf': crear_pdf, '.xlsx': crear_xlsx, '.pptx': crear_pptx}.get(ext)
        if not f: sys.exit('La salida debe terminar en .docx, .xlsx, .pptx o .pdf')
        r = f(S)
        print(json.dumps({'ok': True, 'archivo': S['salida'], 'bytes': os.path.getsize(S['salida']), **r}, ensure_ascii=False))
    else: sys.exit(__doc__)
